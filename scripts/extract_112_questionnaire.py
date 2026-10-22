import openpyxl
import re
import json
import os

wb_path = r"C:\Users\footb\Downloads\NoCo Property Acquisition & Realtor Due Diligence Questionnaire.xlsx"
wb = openpyxl.load_workbook(wb_path)
sheet = wb.active

def clean_text(s):
    if not s:
        return ""
    s = str(s).strip()
    # Replace digits\ufffddigits with digits - digits
    s = re.sub(r'(\d+)\s*\ufffd\s*(\d+)', r'\1–\2', s)
    s = re.sub(r'(\d+%?)\s*\ufffd\s*(\d+%?)', r'\1–\2', s)
    # Replace word\ufffds with word's
    s = re.sub(r'(\w+)\ufffds', r"\1's", s)
    # Replace \ufffd...\ufffd with quotes
    s = re.sub(r'\ufffd([^\ufffd]+)\ufffd', r'"\1"', s)
    # Remaining \ufffd to en-dash
    s = s.replace('\ufffd', '–')
    return s.strip()

data = []
for r in range(8, sheet.max_row + 1):
    q_num = sheet.cell(r, 2).value
    category = sheet.cell(r, 3).value
    q_text = sheet.cell(r, 4).value
    target = sheet.cell(r, 5).value
    source = sheet.cell(r, 6).value
    assessment = sheet.cell(r, 7).value or 'Pending'
    realtor_resp = sheet.cell(r, 8).value or ''
    notes = sheet.cell(r, 9).value or ''
    
    if q_num is not None and category and q_text:
        try:
            q_int = int(float(str(q_num)))
        except:
            q_int = q_num
        data.append({
            'id': f'Q{q_int}',
            'num': q_int,
            'category': clean_text(category),
            'question': clean_text(q_text),
            'benchmark': clean_text(target),
            'source': clean_text(source),
            'assessment': clean_text(assessment),
            'realtorResponse': clean_text(realtor_resp),
            'notes': clean_text(notes)
        })

print(f"Extracted and cleaned {len(data)} due diligence questions.")

target_json = r"C:\AI-BS\frontend\src\components\noco_due_diligence_112.json"
with open(target_json, "w", encoding="utf-8") as f:
    json.dump(data, f, indent=2, ensure_ascii=False)

print(f"Saved cleanly to {target_json}")

# Also output categories summary
cats = {}
for item in data:
    cat = item['category']
    cats[cat] = cats.get(cat, 0) + 1

for c, count in cats.items():
    print(f"  {c}: {count} questions")
